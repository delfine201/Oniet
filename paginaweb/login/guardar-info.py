from pathlib import Path
import openpyxl
import time
from datetime import datetime
from flask import Flask, render_template, jsonify

app = Flask(__name__)

ruta = Path("D://Datos de aucam//Documents//programas//proyectos_vs//Login//usuarios.xlsx")

if ruta.exists():
    wb = openpyxl.load_workbook(ruta)
else:
    wb = openpyxl.Workbook()
    wb.save(ruta)

ws = wb["Sheet"]

ws["A1"] = "ID"
ws["B1"] = "Nombres"
ws["C1"] = "Apellidos"
ws["D1"] = "Nacimiento"
ws["E1"] = "Contraseña"

print(wb.sheetnames)

wb.save(ruta)